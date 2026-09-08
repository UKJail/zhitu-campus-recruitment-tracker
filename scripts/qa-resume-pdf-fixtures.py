"""Create deterministic, wholly synthetic DOCX renderer acceptance fixtures.

Uses the bundled python-docx runtime. Never reads candidate files or credentials.
Binary DOCX files are QA intermediates; a small JSON/base64 bundle is optional
for reproducible transfer to the isolated renderer host without Python packages.
"""
from __future__ import annotations

import argparse
import base64
import hashlib
import io
import json
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from lxml import etree

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
CT = "http://schemas.openxmlformats.org/package/2006/content-types"
KEEP_STYLES = {"Normal", "Title", "Heading1", "Heading2", "Header", "Footer", "DefaultParagraphFont", "TableNormal"}


def font(run, size=11, bold=False):
    run.font.name = "Noto Sans CJK SC"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for kind in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + kind), "Noto Sans CJK SC")


def paragraph(container, text, style=None, size=11, bold=False):
    item = container.add_paragraph(style=style)
    item.paragraph_format.space_after = Pt(5)
    item.paragraph_format.line_spacing = 1.12
    font(item.add_run(text), size, bold)
    return item


def heading(container, text):
    item = paragraph(container, text, "Heading 1", 12, True)
    item.paragraph_format.space_before = Pt(8)
    item.paragraph_format.keep_with_next = True
    return item


def document(title, letter=False):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(215.9 if letter else 210)
    section.page_height = Mm(279.4 if letter else 297)
    section.top_margin = section.bottom_margin = Mm(17)
    section.left_margin = section.right_margin = Mm(18)
    section.header_distance = section.footer_distance = Mm(8)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Header", "Footer"):
        style = doc.styles[name]
        style.font.name = "Noto Sans CJK SC"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(11)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Noto Sans CJK SC")
    doc.core_properties.author = "Zhitu Synthetic QA"
    doc.core_properties.last_modified_by = "Zhitu Synthetic QA"
    doc.core_properties.title = title
    doc.core_properties.subject = "Synthetic resume PDF acceptance fixture"
    doc.core_properties.created = doc.core_properties.modified = datetime(2026, 1, 1, tzinfo=timezone.utc)
    paragraph(doc, title, "Title", 20, True)
    paragraph(doc, "测试候选人甲  Synthetic Candidate Alpha", size=12)
    paragraph(doc, "qa-resume@example.invalid  |  示例城市  |  本文件仅用于软件验收", size=10)
    return doc


def body(container):
    heading(container, "教育经历")
    paragraph(container, "示例大学  商业分析本科  2022.09 - 2026.06")
    paragraph(container, "相关课程包括统计学、财务分析和数据库基础。课程材料仅为本次软件测试编写，不对应真实院校或个人。")
    heading(container, "实习经历")
    paragraph(container, "示例研究团队  数据整理实习生  2025.07 - 2025.09")
    paragraph(container, "• 协助整理公开数据，使用 Excel 检查日期格式和重复记录，形成字段说明文档。")
    paragraph(container, "• 参与周度讨论，记录分析假设与待核对事项，没有独立负责业务或编造商业成果。")
    heading(container, "课程项目")
    paragraph(container, "校园活动数据课程项目  Course Project  2025.03 - 2025.05")
    paragraph(container, "• 使用 SQL 完成分组统计，核对结果与原始数据的一致性，并解释缺失值处理方式。")
    paragraph(container, "• 使用 Python 整理图表，与组员共同完成课程报告；团队结论不表述为个人独立成果。")
    heading(container, "工具与语言")
    paragraph(container, "Excel  SQL  Python  普通话  English reading and writing")
    paragraph(container, "日期重复核验 2025.07 - 2025.09  编号 QA 2026  联系地址 qa-resume@example.invalid")


def compact_package(doc):
    stream = io.BytesIO()
    doc.save(stream)
    with zipfile.ZipFile(stream) as source:
        names = {name for name in source.namelist() if name in {
            "[Content_Types].xml", "_rels/.rels", "docProps/core.xml", "docProps/app.xml",
            "word/document.xml", "word/_rels/document.xml.rels", "word/styles.xml",
        } or name.startswith("word/header") or name.startswith("word/footer")}
        entries = {name: source.read(name) for name in names}
    styles = etree.fromstring(entries["word/styles.xml"])
    for item in list(styles):
        if item.tag == f"{{{W}}}latentStyles" or (item.tag == f"{{{W}}}style" and item.get(f"{{{W}}}styleId") not in KEEP_STYLES):
            styles.remove(item)
    entries["word/styles.xml"] = etree.tostring(styles, xml_declaration=True, encoding="UTF-8", standalone=True)
    root_relationships = etree.fromstring(entries["_rels/.rels"])
    for item in list(root_relationships):
        if item.get("Target", "") not in names:
            root_relationships.remove(item)
    entries["_rels/.rels"] = etree.tostring(root_relationships, xml_declaration=True, encoding="UTF-8", standalone=True)
    relationships = etree.fromstring(entries["word/_rels/document.xml.rels"])
    for item in list(relationships):
        if "word/" + item.get("Target", "") not in names:
            relationships.remove(item)
    entries["word/_rels/document.xml.rels"] = etree.tostring(relationships, xml_declaration=True, encoding="UTF-8", standalone=True)
    types = etree.fromstring(entries["[Content_Types].xml"])
    for item in list(types):
        if item.tag == f"{{{CT}}}Override" and item.get("PartName", "").lstrip("/") not in names:
            types.remove(item)
    entries["[Content_Types].xml"] = etree.tostring(types, xml_declaration=True, encoding="UTF-8", standalone=True)
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for name in sorted(entries):
            info = zipfile.ZipInfo(name, date_time=(2026, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(info, entries[name])
    return output.getvalue()


def fixtures():
    single = document("中文单栏简历验收")
    body(single)
    yield "a4-single-column", None, single

    double = document("中文双栏简历验收")
    table = double.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Mm(57)
    table.columns[1].width = Mm(117)
    for index, cell in enumerate(table.rows[0].cells):
        cell.width = Mm(57 if index == 0 else 117)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        properties = cell._tc.get_or_add_tcPr()
        borders = OxmlElement("w:tcBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            border = OxmlElement("w:" + edge)
            border.set(qn("w:val"), "single")
            border.set(qn("w:sz"), "4")
            border.set(qn("w:color"), "D9D9D9")
            borders.append(border)
        properties.append(borders)
        margins = OxmlElement("w:tcMar")
        for edge in ("top", "left", "bottom", "right"):
            margin = OxmlElement("w:" + edge)
            margin.set(qn("w:w"), "120")
            margin.set(qn("w:type"), "dxa")
            margins.append(margin)
        properties.append(margins)
    left, right = table.rows[0].cells
    heading(left, "教育经历")
    paragraph(left, "示例大学 商业分析本科")
    paragraph(left, "2022.09 - 2026.06")
    paragraph(left, "Statistics and Finance")
    heading(left, "工具与语言")
    paragraph(left, "Excel SQL Python")
    paragraph(left, "普通话 英语读写")
    heading(left, "联系信息")
    paragraph(left, "qa-resume@example.invalid", size=10)
    heading(right, "实习经历")
    paragraph(right, "示例研究团队 数据整理实习生")
    paragraph(right, "2025.07 - 2025.09")
    paragraph(right, "协助整理公开数据，核对日期格式并记录缺失字段，完成可复查的字段说明。")
    paragraph(right, "参与团队讨论，解释清洗步骤与检查范围，不将团队成果归为个人独立成果。")
    heading(right, "课程项目")
    paragraph(right, "Campus Data Course Project")
    paragraph(right, "2025.03 - 2025.05")
    paragraph(right, "使用 SQL 统计活动记录，并以 Python 生成图表。记录假设与数据范围，保留中英文混合说明。")
    paragraph(right, "这一段故意较长，用于检查双栏文字换行后是否被交叉读取。所有语句都应留在右侧栏目，左侧学校与工具不得混入项目描述。")
    yield "a4-double-column-table", None, double

    header = document("页眉页脚文字验收")
    body(header)
    font(header.sections[0].header.paragraphs[0].add_run("中文页眉 HEADER QA 2026"), 9)
    font(header.sections[0].footer.paragraphs[0].add_run("中文页脚 FOOTER QA 2026"), 9)
    yield "a4-header-footer", None, header

    overflow = document("超页拒绝验收")
    body(overflow)
    overflow.add_page_break()
    heading(overflow, "第二页必须保留")
    paragraph(overflow, "该页是明确加入的边界样例。系统应返回 PAGE_COUNT，不得截断、删除文字或自动缩小字号。")
    yield "a4-overflow-two-pages", "PAGE_COUNT", overflow

    letter = document("非 A4 纸张拒绝验收", letter=True)
    heading(letter, "教育经历")
    paragraph(letter, "示例大学 商业分析本科 2022.09 - 2026.06")
    paragraph(letter, "此文件采用 Letter 纸张，预期返回 PAGE_SIZE。姓名、文字与纸张设置都不得由转换服务擅自修改。")
    yield "letter-page-size", "PAGE_SIZE", letter


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--outdir", type=Path, required=True)
    parser.add_argument("--bundle", type=Path)
    args = parser.parse_args()
    args.outdir.mkdir(parents=True, exist_ok=True)
    records = []
    for name, expected, doc in fixtures():
        payload = compact_package(doc)
        (args.outdir / (name + ".docx")).write_bytes(payload)
        records.append({"id": name, "expectedError": expected, "sha256": hashlib.sha256(payload).hexdigest(), "docxBase64": base64.b64encode(payload).decode("ascii")})
    bundle = {"schemaVersion": 1, "syntheticOnly": True, "fixtureCount": len(records), "fixtures": records}
    (args.outdir / "manifest.json").write_text(json.dumps(bundle, ensure_ascii=False, indent=2), encoding="utf-8")
    if args.bundle:
        args.bundle.parent.mkdir(parents=True, exist_ok=True)
        args.bundle.write_text(json.dumps(bundle, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"createdFixtures": len(records), "totalDocxBytes": sum(len(base64.b64decode(item["docxBase64"])) for item in records), "syntheticOnly": True}))


if __name__ == "__main__":
    main()
